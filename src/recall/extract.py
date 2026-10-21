from __future__ import annotations

import logging
from pathlib import Path

from .config import KIND_BY_EXT
from .util import read_text_smart

log = logging.getLogger("recall.extract")

try:
    import pymupdf as fitz  # PyMuPDF（1.x 起推荐模块名，避免 fitz 弃用告警）
except ImportError:
    try:
        import fitz  # type: ignore[no-redef]
    except ImportError:
        fitz = None  # type: ignore[assignment]


def kind_for(path: Path) -> str | None:
    return KIND_BY_EXT.get(path.suffix.lower())


def extract_text(path: Path, max_chars: int) -> tuple[str, str | None]:
    """返回 (文本, 备注)。备注 None 表示正常；失败返回空串。"""
    kind = kind_for(path)
    if kind == "pdf":
        return _extract_pdf(path, max_chars)
    if kind == "docx":
        return _extract_docx(path, max_chars)
    if kind in ("text", "code"):
        return read_text_smart(path)[:max_chars], None
    return "", "unsupported"


def _extract_pdf(path: Path, max_chars: int) -> tuple[str, str | None]:
    if fitz is None:
        log.warning("未安装 PyMuPDF，跳过 PDF：%s", path)
        return "", "no-pymupdf"
    try:
        with fitz.open(path) as doc:
            text = "\n".join(page.get_text("text") for page in doc)
    except Exception as e:
        log.warning("PDF 解析失败 %s：%s", path, e)
        return "", "parse-error"
    return text[:max_chars], None


def _extract_docx(path: Path, max_chars: int) -> tuple[str, str | None]:
    # 主路径 PyMuPDF；抽不到文本（或引擎不支持该 docx）时用 python-docx 兜底，段落+表格更稳
    if fitz is not None:
        try:
            with fitz.open(path, filetype="docx") as doc:
                text = "\n".join(page.get_text("text") for page in doc)
            if text.strip():
                return text[:max_chars], None
        except Exception:
            pass
    try:
        import docx as pydocx
    except ImportError:
        log.warning("docx 抽取失败（无可用引擎）：%s", path)
        return "", "no-docx-engine"
    try:
        d = pydocx.Document(str(path))
        paras = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                paras.append("\t".join(cell.text for cell in row.cells))
        return "\n".join(paras)[:max_chars], None
    except Exception as e:
        log.warning("docx 解析失败 %s：%s", path, e)
        return "", "parse-error"
